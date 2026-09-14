from io import BytesIO
import json
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook
from rowspect import load_table, apply_type_conversions, dataframe_to_csv, dataframe_to_xlsx, compare_dataframes, failing_rows, validate_dataframe
from rowspect.cli import run


def test_csv_preserves_identifiers_codes_and_large_nullable_integers():
    df = load_table(b"id,code,large\n00123,NA,9007199254740993\n00456,NULL,\n", "x.csv")
    assert df.id.tolist() == ["00123", "00456"]
    assert df.code.tolist() == ["NA", "NULL"]
    assert df.large.iloc[0] == 9007199254740993
    assert pd.isna(df.large.iloc[1])
    assert load_table(b"id\n123\n", "x.csv", preserve_text=True).iloc[0, 0] == "123"


def test_xlsx_preserves_literal_text_and_native_numbers():
    book = Workbook()
    book.active.append(["id", "code", "amount"])
    book.active.append(["00123", "NA", 12])
    book.active.append(["123", "NULL", 13])
    buf = BytesIO(); book.save(buf)
    df = load_table(buf.getvalue(), "x.xlsx")
    assert df.id.tolist() == ["00123", "123"]
    assert df.code.tolist() == ["NA", "NULL"]
    assert df.amount.tolist() == [12, 13]


@pytest.mark.parametrize("value", ["9223372036854775808", "-9223372036854775809", "inf", "-inf", "1.0000000000001", "1e999"])
def test_strict_integer_blocks_without_raising_or_mutating(value):
    original = pd.DataFrame({"id": [value, None]})
    out, result = apply_type_conversions(original, [{"column": "id", "target_type": "integer"}])
    assert not result[0]["applied"]
    pd.testing.assert_frame_equal(out, original)


def test_nullable_exact_integer_and_unsafe_float_conversion():
    source = pd.DataFrame({"id": ["9007199254740993", None]})
    out, results = apply_type_conversions(source, [{"column": "id", "target_type": "integer"}])
    assert results[0]["applied"] and out.id.iloc[0] == 9007199254740993
    _, results = apply_type_conversions(source, [{"column": "id", "target_type": "float"}])
    assert not results[0]["applied"]


def test_export_headers_cannot_be_formulas_and_source_is_unchanged():
    source = pd.DataFrame([["=2+2"]], columns=["=1+1"])
    csv = dataframe_to_csv(source).decode()
    assert csv.startswith("'=1+1")
    book = load_workbook(BytesIO(dataframe_to_xlsx(source)), data_only=False)
    assert book.active['A1'].data_type == 's'
    assert book.active['A2'].data_type == 's'
    assert source.columns.tolist() == ["=1+1"]
    assert source.iloc[0, 0] == "=2+2"


def test_cross_column_rules_exact_numeric_and_dates_with_complete_repair_rows():
    source = pd.DataFrame({"refund": ["9007199254740993", "2", None], "total": ["9007199254740992", "3", "4"],
                           "shipped": ["2026-01-01", "2026-01-03", "bad"], "ordered": ["2026-01-02"]*3})
    rules = [{"id": "refund", "type": "compare_columns", "column": "refund", "other_column": "total", "operator": "le"},
             {"id": "dates", "type": "compare_columns", "column": "shipped", "other_column": "ordered", "operator": "ge", "comparison": "date"}]
    result = validate_dataframe(source, rules)
    assert result['violation_count'] == 3
    failed = failing_rows(source, rules)
    assert failed.rowspect_source_row.tolist() == [2, 4]
    assert failed.rowspect_rule_ids.tolist() == ['refund; dates', 'dates']
    bad = [{**rules[0], "other_column": "missing"}]
    assert validate_dataframe(source, bad)['configuration_error_count'] == 1
    with pytest.raises(ValueError):
        failing_rows(source, bad)
    many = pd.DataFrame({"id": [None]*150})
    assert len(failing_rows(many, [{"type": "required", "column": "id"}])) == 150


def test_comparison_detects_schema_missingness_and_new_categories():
    old = pd.DataFrame({"amount": [10, 20], "state": ["AL", "GA"], "removed": [1, 2]})
    new = pd.DataFrame({"amount": [30, None, 50], "state": ["AL", "FL", "GA"], "added": [1, 2, 3]})
    comparison = compare_dataframes(old, new)
    assert comparison['added_columns'] == ['added'] and comparison['removed_columns'] == ['removed']
    assert comparison['row_change'] == 1
    assert comparison['columns'][0]['missing_percentage_point_change'] == pytest.approx(33.333)
    assert comparison['columns'][1]['new_categories'] == ['FL']
    json.dumps(comparison, allow_nan=False)
    with pytest.raises(ValueError, match='duplicate'):
        compare_dataframes(pd.DataFrame([[1, 2]], columns=['x','x']), new)


def test_cli_comparison_and_failing_rows_do_not_overwrite_source(tmp_path):
    current, baseline = tmp_path/'new.csv', tmp_path/'old.csv'
    current.write_text('id,refund,total\n001,5,4\n002,1,2\n')
    baseline.write_text('id,refund,total\n001,2,4\n')
    rules = tmp_path/'rules.json'
    rules.write_text(json.dumps({'rules':[{'type':'compare_columns','column':'refund','other_column':'total','operator':'le'}]}))
    assert run([str(current), '--baseline', str(baseline), '--comparison-json', str(tmp_path/'comparison.json'), '--rules-profile', str(rules), '--failing-rows', str(tmp_path/'failed.csv')]) == 0
    assert '001' in (tmp_path/'failed.csv').read_text()
    original = current.read_bytes()
    assert run([str(current), '--json', str(current)]) == 2
    assert current.read_bytes() == original


def test_cross_column_dates_accept_native_python_date_cells():
    from datetime import date
    from rowspect.validation import validate_dataframe
    frame = pd.DataFrame({'start': [date(2026, 9, 1), date(2026, 9, 4)],
                          'end': [date(2026, 9, 2), date(2026, 9, 3)]})
    result = validate_dataframe(frame, [{'id': 'dates', 'type': 'compare_columns',
        'column': 'start', 'other_column': 'end', 'operator': 'le', 'comparison': 'date'}])
    assert result['results'][0]['violation_count'] == 1
    assert result['results'][0]['row_numbers'] == [3]
