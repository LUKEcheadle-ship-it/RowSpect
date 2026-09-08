from io import BytesIO

import pandas as pd
import pytest

from rowspect.export import dataframe_to_csv, dataframe_to_xlsx
from rowspect.io import RowSpectIOError, load_table, validate_upload
from rowspect.profile import profile_dataframe


@pytest.mark.parametrize(
    ("name", "payload"),
    [
        ("semicolon", b"a;b\n1;x\n"),
        ("tab", b"a\tb\n1\tx\n"),
        ("utf8-bom", b"\xef\xbb\xbfa,b\n1,x\n"),
        ("latin1", "name\nJosé\n".encode("latin-1")),
    ],
)
def test_supported_csv_variants_load_cleanly(name, payload):
    df = load_table(payload, f"{name}.csv")
    expected_shape = (1, 1) if name == "latin1" else (1, 2)
    assert df.shape == expected_shape


@pytest.mark.parametrize(
    ("payload", "filename", "message"),
    [
        (b"", "empty.csv", "empty"),
        (b"a,b\n1,2,3\n", "bad.csv", "inconsistent"),
        (b"a,b\n\"1,2\n", "bad.csv", "malformed quoting"),
        (b"anything", "notes.txt", "supports CSV"),
    ],
)
def test_malformed_or_unsupported_inputs_have_clean_errors(payload, filename, message):
    with pytest.raises(RowSpectIOError, match=message):
        load_table(payload, filename)


def test_upload_size_limit_is_enforced():
    with pytest.raises(RowSpectIOError, match="larger"):
        validate_upload(b"12345", "oversized.csv", max_bytes=4)


def test_blank_headers_remain_visible_to_the_profiler():
    df = load_table(b",\n1,2\n", "blank-headers.csv")
    profile = profile_dataframe(df)
    assert list(df.columns) == ["", ""]
    assert profile["duplicate_columns"] == 1


def test_corrupt_xlsx_has_a_user_facing_error():
    with pytest.raises(RowSpectIOError, match="corrupt|malformed|valid Excel"):
        load_table(b"not-an-xlsx", "broken.xlsx")


def test_all_formula_prefixes_are_neutralized_in_both_exports():
    source = pd.DataFrame({"text": ["=one", "+two", "-three", "@four"]})
    csv = dataframe_to_csv(source).decode("utf-8")
    assert all("'" + value in csv for value in source["text"])
    xlsx = dataframe_to_xlsx(source)
    assert xlsx[:2] == b"PK"
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(xlsx), read_only=True, data_only=False)
    assert [row[0] for row in workbook.active.iter_rows(min_row=2, values_only=True)] == [
        "'=one",
        "'+two",
        "'-three",
        "'@four",
    ]
