"""Exercise real upload/download workflows against an isolated loopback app.

Requires Playwright and a Chromium installation. Uses generated synthetic inputs
only; writes evidence to a new directory outside the checkout. This is automated
browser evidence, not a claim of a human Microsoft Excel walkthrough.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import socket
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

ROOT = Path(__file__).resolve().parents[1]


def run(output: Path, executable: str | None) -> dict:
    output = output.resolve()
    if output.is_relative_to(ROOT):
        raise ValueError("Evidence must be outside the checkout")
    output.mkdir(parents=True, exist_ok=False)
    inputs, downloads = output / "inputs", output / "downloads"
    inputs.mkdir(); downloads.mkdir()
    valid = pd.DataFrame({"id": ["001", "002", "002", "003"],
                          "amount": ["10", "20", "20", "30"],
                          "note": ["=1+1", " NA ", " NA ", "NULL"]})
    invalid = valid.copy(); invalid.loc[3, "amount"] = "bad-number"
    invalid.to_csv(inputs / "invalid.csv", index=False)
    with pd.ExcelWriter(inputs / "current.xlsx", engine="openpyxl") as writer:
        valid.iloc[:1].to_excel(writer, sheet_name="Other", index=False)
        valid.to_excel(writer, sheet_name="Current", index=False)
        for sheet in writer.sheets.values():
            sheet.cell(2, 3).data_type = "s"
    with pd.ExcelWriter(inputs / "baseline.xlsx", engine="openpyxl") as writer:
        valid.iloc[:1].to_excel(writer, sheet_name="Older", index=False)
        valid.iloc[:2].to_excel(writer, sheet_name="0", index=False)
        for sheet in writer.sheets.values():
            sheet.cell(2, 3).data_type = "s"
    profile = {"version": 1, "name": "Synthetic review", "description": "Browser qualification",
               "rules": [{"id": "required-id", "type": "required", "column": "id"},
                         {"id": "unique-id", "type": "unique", "column": "id"}],
               "conversions": [{"id": "amount-int", "column": "amount", "target_type": "integer"}]}
    (inputs / "rules.json").write_text(json.dumps(profile), encoding="utf-8")
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs.iterdir()}
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0)); port = sock.getsockname()[1]
    checks = []
    with (output / "server.log").open("w", encoding="utf-8") as log:
        process = subprocess.Popen([sys.executable, "-m", "streamlit", "run", str(ROOT / "app.py"),
            "--server.headless=true", "--server.address=127.0.0.1", f"--server.port={port}",
            "--browser.gatherUsageStats=false"], cwd=ROOT, stdout=log, stderr=subprocess.STDOUT)
        try:
            deadline = time.monotonic() + 30
            while True:
                if process.poll() is not None:
                    raise RuntimeError("Streamlit exited before readiness")
                try:
                    with urllib.request.urlopen(f"http://127.0.0.1:{port}/_stcore/health", timeout=1) as response:
                        if response.status == 200: break
                except OSError:
                    if time.monotonic() >= deadline: raise RuntimeError("Streamlit readiness timed out")
                    time.sleep(.2)
            with sync_playwright() as playwright:
                kwargs = {"headless": True, "args": ["--disable-dev-shm-usage"]}
                if executable: kwargs["executable_path"] = executable
                browser = playwright.chromium.launch(**kwargs)
                try:
                    context = browser.new_context(accept_downloads=True, viewport={"width": 1450, "height": 1050})
                    page = context.new_page(); page.set_default_timeout(15000)
                    def settled():
                        # Observe the real frontend lifecycle, not an arbitrary sleep
                        # or a replacement data-loading path. Require a stable idle
                        # interval so a queued rerun cannot discard the next upload.
                        page.evaluate("window.__reviewIdleSince = null")
                        page.wait_for_function("""() => {
                            const app = document.querySelector('[data-testid="stApp"]');
                            const idle = app && app.dataset.testConnectionState === 'CONNECTED'
                                && app.dataset.testScriptState === 'notRunning';
                            if (!idle) { window.__reviewIdleSince = null; return false; }
                            if (window.__reviewIdleSince === null) window.__reviewIdleSince = performance.now();
                            return performance.now() - window.__reviewIdleSince >= 600;
                        }""", timeout=30000)
                        expect(page.get_by_test_id("stException")).to_have_count(0)
                    def upload(label, file):
                        settled()
                        widget = page.get_by_test_id("stFileUploader").filter(
                            has=page.get_by_text(label, exact=True))
                        widget.locator('input[type="file"]').set_input_files(file)
                        expect(widget.get_by_text(file.name, exact=True)).to_be_visible()
                        settled()
                        expect(widget.get_by_text(file.name, exact=True)).to_be_visible()
                    def tab(name):
                        settled()
                        page.get_by_role("tab", name=name, exact=True).click()
                        expect(page.get_by_test_id("stException")).to_have_count(0)
                    def select(label, value):
                        widget = page.get_by_test_id("stSelectbox").filter(has=page.get_by_text(label, exact=True))
                        widget.get_by_role("combobox").click()
                        page.get_by_role("option", name=value, exact=True).click()
                        expect(widget.get_by_role("combobox")).to_have_value(value)
                        settled()
                    def download(label, name):
                        with page.expect_download() as info:
                            page.get_by_role("button", name=label, exact=True).click()
                        target = downloads / name; info.value.save_as(target)
                        assert target.stat().st_size > 0
                        return target
                    page.goto(f"http://127.0.0.1:{port}")
                    page.get_by_text("Use the built-in sample", exact=True).click()
                    expect(page.get_by_label("Use the built-in sample", exact=True)).to_be_checked()
                    settled()
                    expect(page.get_by_role("tab", name="Overview", exact=True)).to_be_visible()
                    for name in ["Overview", "Issues", "Columns", "Explore", "Compare", "Validate", "Convert", "Clean", "Export"]:
                        tab(name)
                    checks.append("built-in sample renders all nine tabs without application exceptions")
                    # Real browser file upload, not a patched widget or direct API call.
                    upload("CSV or Excel", inputs / "invalid.csv")
                    tab("Validate")
                    upload("Load a RowSpect rule profile", inputs / "rules.json")
                    page.get_by_role("button", name="Use loaded profile in this session", exact=True).click()
                    expect(page.get_by_text("This dataset passes every configured validation rule.", exact=True)).to_have_count(0)
                    download("Download validation results JSON", "validation.json")
                    download("Download failing rows CSV", "failing-rows.csv")
                    download("Download reusable rule profile", "profile.json")
                    tab("Convert")
                    expect(page.get_by_text("At least one planned conversion is blocked.", exact=False)).to_be_visible()
                    expect(page.get_by_role("button", name="Download converted CSV", exact=True)).to_have_count(0)
                    checks.append("CSV upload, rule-profile load, duplicate validation, failed-row export and unsafe conversion refusal")
                    upload("CSV or Excel", inputs / "current.xlsx")
                    select("Worksheet", "Current")
                    expect(page.get_by_text("Source: current.xlsx", exact=False)).to_contain_text("sheet: Current")
                    settled()
                    upload("Previous baseline (optional)", inputs / "baseline.xlsx")
                    select("Baseline worksheet", "0")
                    tab("Compare")
                    compare = json.loads(download("Download comparison JSON", "comparison.json").read_text())
                    assert compare["row_count_change"] == 2, compare
                    assert not page.get_by_text("Baseline could not be compared", exact=False).count()
                    checks.append("independent XLSX worksheets, numeric sheet name 0 and correct baseline comparison")
                    tab("Convert")
                    expect(page.get_by_role("button", name="Download converted CSV", exact=True)).to_be_visible()
                    download("Download converted CSV", "converted.csv")
                    download("Download converted Excel", "converted.xlsx")
                    tab("Clean")
                    download("Download cleaned CSV", "cleaned.csv")
                    download("Download cleaned Excel", "cleaned.xlsx")
                    tab("Export")
                    download("Download HTML report", "report.html")
                    download("Download JSON profile", "analysis.json")
                    checks.append("safe conversion, cleanup, CSV/XLSX/HTML/JSON browser downloads")
                    for filename, rows in [("converted.csv", 4), ("cleaned.csv", 3)]:
                        frame = pd.read_csv(downloads / filename, keep_default_na=False, dtype=str)
                        assert len(frame) == rows and frame.iloc[0]["id"] == "001"
                        assert frame.iloc[0]["note"].startswith("'=")
                    for filename, rows in [("converted.xlsx", 4), ("cleaned.xlsx", 3)]:
                        book = load_workbook(downloads / filename, data_only=False)
                        try:
                            sheet = book.active
                            assert sheet.max_row == rows + 1
                            assert sheet["A2"].value == "001"
                            assert sheet["C2"].data_type != "f"
                            assert str(sheet["C2"].value).startswith("'=")
                        finally: book.close()
                    assert json.loads((downloads / "profile.json").read_text())["conversions"][0]["target_type"] == "integer"
                    assert json.loads((downloads / "validation.json").read_text())["validation_passed"] is False
                    assert before == {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in inputs.iterdir()}
                    checks.append("download contents verified: row counts, padded identifiers, formula neutralization and input immutability")
                    tab("Overview");page.screenshot(path=str(output / "overview.png"), full_page=True)
                    tab("Compare");page.screenshot(path=str(output / "comparison.png"), full_page=True)
                    tab("Validate");page.screenshot(path=str(output / "validation.png"), full_page=True)
                    # Reload exported profile through a fresh browser session.
                    context.close(); context = browser.new_context(accept_downloads=True)
                    page = context.new_page();page.set_default_timeout(15000);page.goto(f"http://127.0.0.1:{port}")
                    upload("CSV or Excel", inputs / "invalid.csv")
                    tab("Validate")
                    upload("Load a RowSpect rule profile", downloads / "profile.json")
                    page.get_by_role("button", name="Use loaded profile in this session", exact=True).click()
                    tab("Convert")
                    expect(page.get_by_text("At least one planned conversion is blocked.", exact=False)).to_be_visible()
                    checks.append("exported reusable profile reloads in a fresh browser session with refusal preserved")
                    context.close()
                except Exception:
                    import traceback
                    (output / "failure-traceback.txt").write_text(traceback.format_exc(), encoding="utf-8")
                    if "page" in locals() and not page.is_closed():
                        try:
                            page.screenshot(path=str(output / "failure.png"), full_page=True)
                            (output / "failure.html").write_text(page.content(), encoding="utf-8")
                            (output / "failure-text.txt").write_text(page.locator("body").inner_text(), encoding="utf-8")
                        except Exception as capture_error:
                            print(f"Failure capture unavailable: {capture_error}", file=sys.stderr)
                    raise
                finally:
                    browser.close()
        finally:
            process.terminate()
            try: process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill();process.wait(timeout=5)
    receipt = {"decision": "PASS", "checks": checks, "synthetic_inputs_only": True,
               "native_spreadsheet_application_checked": False,
               "source_sha": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT,text=True).strip(),
               "source_dirty": bool(subprocess.check_output(["git", "status", "--porcelain"],cwd=ROOT,text=True)),
               "app_sha256": hashlib.sha256((ROOT / "app.py").read_bytes()).hexdigest(),
               "files_sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in downloads.iterdir()}}
    (output / "receipt.json").write_text(json.dumps(receipt,indent=2)+"\n",encoding="utf-8")
    print(json.dumps(receipt,indent=2));return receipt


if __name__ == "__main__":
    parser=argparse.ArgumentParser()
    parser.add_argument("--output-dir",required=True,type=Path)
    parser.add_argument("--browser-executable")
    args=parser.parse_args();run(args.output_dir,args.browser_executable)
